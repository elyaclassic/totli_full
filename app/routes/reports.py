"""
Hisobotlar — savdo, qoldiq, qarzdorlik va Excel export.
"""
import io
from datetime import datetime
from fastapi import APIRouter, Request, Depends, File, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.core import templates
from app.models.database import get_db, Order, Stock, Product, Partner, Warehouse, User
from app.deps import get_current_user, require_auth

router = APIRouter(prefix="/reports", tags=["reports"])


@router.get("", response_class=HTMLResponse)
async def reports_index(request: Request, current_user: User = Depends(require_auth)):
    """Hisobotlar bosh sahifasi"""
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse("reports/index.html", {
        "request": request,
        "page_title": "Hisobotlar",
        "current_user": current_user,
    })


@router.get("/sales", response_class=HTMLResponse)
async def report_sales(
    request: Request,
    start_date: str = None,
    end_date: str = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_auth),
):
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    if not start_date:
        start_date = datetime.now().replace(day=1).strftime("%Y-%m-%d")
    if not end_date:
        end_date = datetime.now().strftime("%Y-%m-%d")
    orders = db.query(Order).filter(
        Order.type == "sale",
        Order.date >= start_date,
        Order.date <= end_date + " 23:59:59",
    ).all()
    total = sum(o.total for o in orders)
    return templates.TemplateResponse("reports/sales.html", {
        "request": request,
        "orders": orders,
        "total": total,
        "start_date": start_date,
        "end_date": end_date,
        "page_title": "Savdo hisoboti",
        "current_user": current_user,
    })


@router.get("/sales/export")
async def report_sales_export(
    start_date: str = None,
    end_date: str = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_auth),
):
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    if not start_date:
        start_date = datetime.now().replace(day=1).strftime("%Y-%m-%d")
    if not end_date:
        end_date = datetime.now().strftime("%Y-%m-%d")
    orders = db.query(Order).filter(
        Order.type == "sale",
        Order.date >= start_date,
        Order.date <= end_date + " 23:59:59",
    ).order_by(Order.date.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Savdo"
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws["A1"] = "Savdo hisoboti"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Davr: {start_date} — {end_date}"
    ws.append(["№", "Sana", "Buyurtma", "Mijoz", "Jami", "Holat"])
    for c in range(1, 7):
        ws.cell(row=4, column=c).fill = header_fill
        ws.cell(row=4, column=c).font = Font(bold=True, color="FFFFFF")
    for i, o in enumerate(orders, 1):
        ws.append([
            i,
            o.date.strftime("%d.%m.%Y %H:%M") if o.date else "",
            o.number or "",
            o.partner.name if o.partner else "",
            float(o.total or 0),
            o.status or "",
        ])
    total = sum(o.total or 0 for o in orders)
    ws.append(["", "", "", "JAMI:", total, ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=savdo_{start_date}_{end_date}.xlsx"},
    )


@router.get("/stock", response_class=HTMLResponse)
async def report_stock(request: Request, db: Session = Depends(get_db), current_user: User = Depends(require_auth)):
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    stocks = db.query(Stock).join(Product).all()
    return templates.TemplateResponse("reports/stock.html", {
        "request": request,
        "stocks": stocks,
        "page_title": "Qoldiq hisoboti",
        "current_user": current_user,
    })


def _stock_import_column_map(ws):
    """
    Locate header + column indices for stock import.
    Andoza/import layout: Ombor, Mahsulot, Qoldiq [, Tannarx, Sotuv].
    Legacy report export put Kod before Qoldiq — detect by header names so
    re-importing an export cannot treat barcodes as quantities.
    Returns (data_start_row_1based, wh_i, prod_i, qty_i, tannarx_i|None, sotuv_i|None).
    """
    default = (2, 0, 1, 2, 3, 4)
    for r in range(1, 6):
        raw = next(ws.iter_rows(min_row=r, max_row=r, values_only=True), None)
        if not raw:
            continue
        headers = [str(c or "").strip().lower() for c in raw]
        if not any(headers):
            continue
        qty_i = next((i for i, h in enumerate(headers) if "qoldiq" in h), None)
        wh_i = next((i for i, h in enumerate(headers) if "ombor" in h), None)
        if qty_i is None or wh_i is None:
            continue
        prod_i = next(
            (i for i, h in enumerate(headers) if "mahsulot" in h or "tovar" in h or h == "nomi"),
            None,
        )
        if prod_i is None:
            prod_i = wh_i + 1 if wh_i + 1 != qty_i else wh_i + 2
        tannarx_i = next((i for i, h in enumerate(headers) if "tannarx" in h), None)
        if tannarx_i is None:
            # Legacy export labeled purchase price as "Narx" (not min stock / summa).
            tannarx_i = next((i for i, h in enumerate(headers) if h == "narx"), None)
        sotuv_i = next((i for i, h in enumerate(headers) if "sotuv" in h), None)
        return (r + 1, wh_i, prod_i, qty_i, tannarx_i, sotuv_i)
    return default


@router.get("/stock/export")
async def report_stock_export(db: Session = Depends(get_db), current_user: User = Depends(require_auth)):
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    stocks = (
        db.query(Stock)
        .join(Product, Stock.product_id == Product.id)
        .join(Warehouse, Stock.warehouse_id == Warehouse.id)
        .order_by(Warehouse.name, Product.name)
        .all()
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Qoldiq"
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws["A1"] = "Qoldiq hisoboti"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = datetime.now().strftime("%d.%m.%Y %H:%M")
    # First five columns match /reports/stock/import (and andoza) so Excel→Import
    # cannot treat Kod/barcode as Qoldiq. Extra report columns follow.
    ws.append([
        "Ombor",
        "Mahsulot",
        "Qoldiq",
        "Tannarx (so'm)",
        "Sotuv narxi (so'm)",
        "Kod",
        "Minimal",
        "Summa",
    ])
    for c in range(1, 9):
        ws.cell(row=4, column=c).fill = header_fill
        ws.cell(row=4, column=c).font = Font(bold=True, color="FFFFFF")
    for s in stocks:
        p = s.product
        wh = s.warehouse
        min_s = getattr(p, "min_stock", 0) or 0
        purchase = getattr(p, "purchase_price", 0) or 0
        sale = getattr(p, "sale_price", 0) or 0
        qty = float(s.quantity or 0)
        ws.append([
            wh.name if wh else "",
            p.name if p else "",
            qty,
            float(purchase),
            float(sale),
            (p.barcode or p.code or "") if p else "",
            float(min_s),
            float(qty * purchase),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=qoldiq_{datetime.now().strftime('%Y%m%d')}.xlsx"},
    )


@router.get("/stock/andoza")
async def report_stock_andoza(current_user: User = Depends(require_auth)):
    """Qoldiqlar uchun Excel andoza (Tannarx va Sotuv narxi ixtiyoriy)."""
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    wb = Workbook()
    ws = wb.active
    ws.title = "Andoza"
    ws.append(["Ombor nomi (yoki kodi)", "Mahsulot nomi (yoki kodi)", "Qoldiq", "Tannarx (so'm)", "Sotuv narxi (so'm)"])
    ws.append(["Xom ashyo ombori", "Yong'oq", 30, "", ""])
    ws.append(["Xom ashyo ombori", "Bodom", 100, "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=qoldiqlar_andoza.xlsx"},
    )


@router.post("/stock/import")
async def report_stock_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_auth),
):
    """Exceldan qoldiqlarni yuklash. Ustunlar: Ombor nomi (yoki kodi), Mahsulot nomi (yoki kodi), Qoldiq; ixtiyoriy: Tannarx, Sotuv narxi."""
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    contents = await file.read()
    wb = load_workbook(io.BytesIO(contents))
    ws = wb.active
    data_start, wh_i, prod_i, qty_i, tannarx_i, sotuv_i = _stock_import_column_map(ws)
    rows = list(ws.iter_rows(min_row=data_start, values_only=True))
    for row in rows:
        if not row:
            continue
        wh_cell = row[wh_i] if len(row) > wh_i else None
        prod_cell = row[prod_i] if len(row) > prod_i else None
        if wh_cell is None or wh_cell == "" or prod_cell is None or prod_cell == "":
            continue
        wh_key = str(wh_cell or "").strip()
        raw_prod = prod_cell
        if raw_prod is not None and isinstance(raw_prod, (int, float)) and float(raw_prod) == int(float(raw_prod)):
            product_key = str(int(float(raw_prod)))
        else:
            product_key = str(raw_prod or "").strip()
        raw_qty = row[qty_i] if len(row) > qty_i else None
        if raw_qty is None or raw_qty == "":
            continue
        try:
            qty = float(raw_qty)
        except (TypeError, ValueError):
            continue
        if qty < 0:
            continue
        tannarx = None
        sotuv_narxi = None
        if tannarx_i is not None and len(row) > tannarx_i and row[tannarx_i] is not None and row[tannarx_i] != "":
            try:
                tannarx = float(row[tannarx_i])
            except (TypeError, ValueError):
                pass
        if sotuv_i is not None and len(row) > sotuv_i and row[sotuv_i] is not None and row[sotuv_i] != "":
            try:
                sotuv_narxi = float(row[sotuv_i])
            except (TypeError, ValueError):
                pass
        wh = db.query(Warehouse).filter(
            (func.lower(Warehouse.name) == wh_key.lower()) | (Warehouse.code == wh_key)
        ).first()
        product = db.query(Product).filter(
            (Product.code == product_key) | (Product.barcode == product_key)
        ).first()
        if not product and product_key:
            product = db.query(Product).filter(
                Product.name.isnot(None),
                func.lower(Product.name) == product_key.lower()
            ).first()
        if not wh or not product:
            continue
        stock = db.query(Stock).filter(
            Stock.warehouse_id == wh.id,
            Stock.product_id == product.id,
        ).first()
        if stock:
            stock.quantity = qty
        else:
            stock = Stock(warehouse_id=wh.id, product_id=product.id, quantity=qty)
            db.add(stock)
        # Only apply positive prices (blank/0 must not wipe catalog costs like warehouse import).
        if tannarx is not None and tannarx > 0:
            product.purchase_price = tannarx
        if sotuv_narxi is not None and sotuv_narxi > 0:
            product.sale_price = sotuv_narxi
        db.commit()
    return RedirectResponse(url="/reports/stock", status_code=303)


@router.get("/debts", response_class=HTMLResponse)
async def report_debts(request: Request, db: Session = Depends(get_db), current_user: User = Depends(require_auth)):
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    debtors = db.query(Partner).filter(Partner.balance != 0).all()
    total_debt = sum(p.balance for p in debtors if p.balance > 0)
    total_credit = sum(abs(p.balance) for p in debtors if p.balance < 0)
    return templates.TemplateResponse("reports/debts.html", {
        "request": request,
        "debtors": debtors,
        "total_debt": total_debt,
        "total_credit": total_credit,
        "page_title": "Qarzdorlik hisoboti",
        "current_user": current_user,
    })


@router.get("/debts/export")
async def report_debts_export(db: Session = Depends(get_db), current_user: User = Depends(require_auth)):
    if not current_user:
        return RedirectResponse(url="/login", status_code=303)
    debtors = db.query(Partner).filter(Partner.balance != 0).order_by(Partner.name).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Qarzdorlik"
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws["A1"] = "Qarzdorlik hisoboti"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = datetime.now().strftime("%d.%m.%Y %H:%M")
    ws.append(["Kod", "Mijoz", "Telefon", "Balans (qarz +)", "Kredit limiti"])
    for c in range(1, 6):
        ws.cell(row=4, column=c).fill = header_fill
        ws.cell(row=4, column=c).font = Font(bold=True, color="FFFFFF")
    for p in debtors:
        ws.append([
            p.code or "",
            p.name or "",
            p.phone or "",
            float(p.balance or 0),
            float(p.credit_limit or 0),
        ])
    total = sum(p.balance for p in debtors if (p.balance or 0) > 0)
    ws.append(["", "", "JAMI QARZ:", total, ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=qarzdorlik_{datetime.now().strftime('%Y%m%d')}.xlsx"},
    )
