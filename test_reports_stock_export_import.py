"""Regression: reports stock Excel export/import column alignment."""
import asyncio
import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.models.database import Base, Product, Stock, User, Warehouse
from app.routes import reports as reports_routes


@pytest.fixture()
def db():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    try:
        yield session
    finally:
        session.close()


def run(coro):
    return asyncio.run(coro)


class _Upload:
    def __init__(self, content: bytes, filename: str = "qoldiq.xlsx"):
        self.file = io.BytesIO(content)
        self.filename = filename

    async def read(self):
        return self.file.getvalue()


def _seed(db, barcode="4607012345678"):
    user = User(
        username="admin",
        password_hash="x",
        full_name="Admin",
        role="admin",
        is_active=True,
    )
    wh = Warehouse(code="WH1", name="Asosiy ombor", is_active=True)
    product = Product(
        code="P100",
        name="Yong'oq",
        type="xom",
        barcode=barcode,
        is_active=True,
        sale_price=5000,
        purchase_price=3000,
        min_stock=10,
    )
    db.add_all([user, wh, product])
    db.flush()
    stock = Stock(warehouse_id=wh.id, product_id=product.id, quantity=50)
    db.add(stock)
    db.commit()
    return user, wh, product, stock


def test_legacy_export_layout_reimport_does_not_invent_from_barcode(db):
    """Old export put Kod before Qoldiq; re-import must use Qoldiq, not barcode."""
    user, wh, product, stock = _seed(db)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Qoldiq hisoboti"
    ws["A2"] = datetime.now().strftime("%d.%m.%Y %H:%M")
    ws.append(["Ombor", "Mahsulot", "Kod", "Qoldiq", "Minimal", "Narx", "Summa"])
    ws.append([wh.name, product.name, product.barcode, 50, 10, 3000, 150000])
    buf = io.BytesIO()
    wb.save(buf)

    resp = run(reports_routes.report_stock_import(_Upload(buf.getvalue()), db=db, current_user=user))
    assert resp.status_code == 303

    db.refresh(stock)
    db.refresh(product)
    assert stock.quantity == 50
    assert product.purchase_price == 3000
    assert product.sale_price == 5000


async def _read_streaming(resp):
    chunks = []
    async for chunk in resp.body_iterator:
        chunks.append(chunk if isinstance(chunk, (bytes, bytearray)) else bytes(chunk))
    return b"".join(chunks)


def test_new_export_roundtrip_preserves_qty_and_prices(db):
    user, wh, product, stock = _seed(db, barcode="999888777666")
    export_resp = run(reports_routes.report_stock_export(db=db, current_user=user))
    content = run(_read_streaming(export_resp))

    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    header_row = None
    for r in range(1, 6):
        vals = [c.value for c in ws[r]]
        if vals and vals[0] == "Ombor" and vals[2] == "Qoldiq":
            header_row = vals
            break
    assert header_row is not None
    assert "Kod" in header_row

    # Mutate on-hand via import of the exported file (same layout).
    stock.quantity = 12
    product.purchase_price = 1111
    product.sale_price = 2222
    db.commit()

    resp = run(reports_routes.report_stock_import(_Upload(content), db=db, current_user=user))
    assert resp.status_code == 303
    db.refresh(stock)
    db.refresh(product)
    assert stock.quantity == 50
    assert product.purchase_price == 3000
    assert product.sale_price == 5000


def test_andoza_style_import_still_works(db):
    user, wh, product, stock = _seed(db, barcode="111")
    wb = Workbook()
    ws = wb.active
    ws.append(["Ombor nomi (yoki kodi)", "Mahsulot nomi (yoki kodi)", "Qoldiq", "Tannarx (so'm)", "Sotuv narxi (so'm)"])
    ws.append([wh.name, product.name, 77, 4000, 6000])
    buf = io.BytesIO()
    wb.save(buf)

    resp = run(reports_routes.report_stock_import(_Upload(buf.getvalue()), db=db, current_user=user))
    assert resp.status_code == 303
    db.refresh(stock)
    db.refresh(product)
    assert stock.quantity == 77
    assert product.purchase_price == 4000
    assert product.sale_price == 6000


def test_zero_tannarx_does_not_wipe_purchase_price(db):
    user, wh, product, stock = _seed(db, barcode="222")
    wb = Workbook()
    ws = wb.active
    ws.append(["Ombor", "Mahsulot", "Qoldiq", "Tannarx", "Sotuv"])
    ws.append([wh.name, product.name, 50, 0, 0])
    buf = io.BytesIO()
    wb.save(buf)

    resp = run(reports_routes.report_stock_import(_Upload(buf.getvalue()), db=db, current_user=user))
    assert resp.status_code == 303
    db.refresh(product)
    assert product.purchase_price == 3000
    assert product.sale_price == 5000
